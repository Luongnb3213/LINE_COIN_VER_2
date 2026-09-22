"""Test cho flows/xlsx_store.py: đọc/ghi cột Phase 2 trên sheet Excel thật.

Dùng workbook openpyxl tạo trong thư mục tạm — không đụng tới file Excel thật
của người dùng (`config.json`'s `xlsx_path`).
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from flows.xlsx_store import GOOGLE_PREPARE_HEADERS, PHASE2_HEADERS, XlsxStore, XlsxStoreError


class XlsxStoreTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.path = Path(self._tmp.name) / "mails.xlsx"

    def _write_workbook(self, headers: list[str], rows: list[list[object]], sheet_name: str = "Mails") -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(self.path)

    def test_missing_file_raises(self) -> None:
        with self.assertRaises(XlsxStoreError):
            XlsxStore(self.path)

    def test_find_email_true_when_present(self) -> None:
        self._write_workbook(["email"], [["user@example.com"]])
        store = XlsxStore(self.path)
        self.assertTrue(store.find_email("user@example.com"))

    def test_find_email_false_when_absent(self) -> None:
        self._write_workbook(["email"], [["user@example.com"]])
        store = XlsxStore(self.path)
        self.assertFalse(store.find_email("nobody@example.com"))

    def test_find_email_matches_pipe_separated_case_insensitively(self) -> None:
        self._write_workbook(["email"], [["First@Example.com|second@example.com"]])
        store = XlsxStore(self.path)
        self.assertTrue(store.find_email("SECOND@EXAMPLE.COM"))

    def test_read_existing_codes_returns_empty_for_unknown_email(self) -> None:
        self._write_workbook(
            ["email", "gift_code", "gift_code_ministop"],
            [["someone@example.com", "ABC", "DEF"]],
        )
        store = XlsxStore(self.path)
        self.assertEqual(store.read_existing_codes("nobody@example.com"), ("", ""))

    def test_read_existing_codes_matches_pipe_separated_emails_case_insensitively(self) -> None:
        self._write_workbook(
            ["email", "gift_code", "gift_code_ministop"],
            [["First@Example.com|second@example.com", "CODE1", "CODE2"]],
        )
        store = XlsxStore(self.path)
        self.assertEqual(store.read_existing_codes("SECOND@EXAMPLE.COM"), ("CODE1", "CODE2"))

    def test_update_phase2_auto_creates_missing_columns(self) -> None:
        self._write_workbook(["email"], [["user@example.com"]])
        store = XlsxStore(self.path)
        store.update_phase2(
            "user@example.com",
            status="SUCCESS",
            message="Đã lấy đủ gift code.",
            gift_code="GRM29G7WLBMZVL3C",
            gift_code_ministop="GY5LSVDB5C73LZNA",
        )

        wb = openpyxl.load_workbook(self.path)
        ws = wb["Mails"]
        header_row = [cell.value for cell in ws[1]]
        for name in PHASE2_HEADERS:
            self.assertIn(name, header_row)

        headers = {name: idx + 1 for idx, name in enumerate(header_row)}
        self.assertEqual(ws.cell(row=2, column=headers["gift_code"]).value, "GRM29G7WLBMZVL3C")
        self.assertEqual(ws.cell(row=2, column=headers["gift_code_ministop"]).value, "GY5LSVDB5C73LZNA")
        self.assertEqual(ws.cell(row=2, column=headers["phase2_status"]).value, "SUCCESS")

    def test_update_phase2_success_clears_old_error_details(self) -> None:
        self._write_workbook(
            ["email", "error_details"],
            [["user@example.com", "Dừng đột ngột"]],
        )
        store = XlsxStore(self.path)
        store.update_phase2(
            "user@example.com",
            status="SUCCESS",
            message="ok",
            gift_code="GRM29G7WLBMZVL3C",
            gift_code_ministop="GY5LSVDB5C73LZNA",
        )

        wb = openpyxl.load_workbook(self.path)
        ws = wb["Mails"]
        headers = {str(c.value): idx + 1 for idx, c in enumerate(ws[1])}
        self.assertIsNone(ws.cell(row=2, column=headers["error_details"]).value)

    def test_update_phase2_partial_does_not_clobber_existing_other_code(self) -> None:
        self._write_workbook(
            ["email", "gift_code", "gift_code_ministop", "phase2_status"],
            [["user@example.com", "", "GY5LSVDB5C73LZNA", "PARTIAL"]],
        )
        store = XlsxStore(self.path)
        #: Lần chạy này chỉ lấy được airwallet — không được ghi đè mất
        #: `gift_code_ministop` đã có từ trước vì tham số truyền vào rỗng.
        store.update_phase2(
            "user@example.com",
            status="SUCCESS",
            message="Đã lấy đủ gift code.",
            gift_code="GRM29G7WLBMZVL3C",
            gift_code_ministop="",
        )

        wb = openpyxl.load_workbook(self.path)
        ws = wb["Mails"]
        headers = {str(c.value): idx + 1 for idx, c in enumerate(ws[1])}
        self.assertEqual(ws.cell(row=2, column=headers["gift_code"]).value, "GRM29G7WLBMZVL3C")
        self.assertEqual(ws.cell(row=2, column=headers["gift_code_ministop"]).value, "GY5LSVDB5C73LZNA")
        self.assertEqual(ws.cell(row=2, column=headers["phase2_status"]).value, "SUCCESS")

    def test_update_phase2_raises_when_email_not_found(self) -> None:
        self._write_workbook(["email"], [["someone@example.com"]])
        store = XlsxStore(self.path)
        with self.assertRaises(XlsxStoreError):
            store.update_phase2("nobody@example.com", status="FAILED", message="x")

    def test_update_phase2_wraps_locked_file_error(self) -> None:
        #: File Excel đang mở trong Excel -> os.replace() báo OSError. Phải
        #: nổi lên thành XlsxStoreError rõ ràng (GUI/CLI đều đang chỉ bắt
        #: XlsxStoreError ở nơi gọi `update_phase2`), không phải OSError thô.
        self._write_workbook(["email"], [["user@example.com"]])
        store = XlsxStore(self.path)
        with mock.patch("flows.xlsx_store.os.replace", side_effect=OSError("locked")):
            with self.assertRaises(XlsxStoreError):
                store.update_phase2("user@example.com", status="FAILED", message="x")

    def test_wrong_sheet_name_raises(self) -> None:
        self._write_workbook(["email"], [["user@example.com"]], sheet_name="Mails")
        store = XlsxStore(self.path, sheet_name="OtherSheet")
        with self.assertRaises(XlsxStoreError):
            store.read_existing_codes("user@example.com")

    def test_update_line_success_sets_status_and_clears_error(self) -> None:
        self._write_workbook(
            ["email", "status", "error_details"],
            [["user@example.com", "PENDING", "Lỗi lần trước"]],
        )
        store = XlsxStore(self.path)
        store.update_line_success("user@example.com")

        wb = openpyxl.load_workbook(self.path)
        ws = wb["Mails"]
        headers = {str(c.value): idx + 1 for idx, c in enumerate(ws[1])}
        self.assertEqual(ws.cell(row=2, column=headers["line_status"]).value, "SUCCESS")
        self.assertEqual(ws.cell(row=2, column=headers["status"]).value, "SUCCESS")
        self.assertIsNone(ws.cell(row=2, column=headers["error_details"]).value)

    def test_update_line_success_raises_when_email_not_found(self) -> None:
        self._write_workbook(["email"], [["someone@example.com"]])
        store = XlsxStore(self.path)
        with self.assertRaises(XlsxStoreError):
            store.update_line_success("nobody@example.com")

    def test_update_line_failed_sets_status_and_error(self) -> None:
        self._write_workbook(["email", "status"], [["user@example.com", "PENDING"]])
        store = XlsxStore(self.path)
        store.update_line_failed("user@example.com", "Không tìm thấy tài khoản Google")

        wb = openpyxl.load_workbook(self.path)
        ws = wb["Mails"]
        headers = {str(c.value): idx + 1 for idx, c in enumerate(ws[1])}
        self.assertEqual(ws.cell(row=2, column=headers["line_status"]).value, "FAILED")
        self.assertEqual(ws.cell(row=2, column=headers["status"]).value, "FAILED")
        self.assertEqual(ws.cell(row=2, column=headers["error_details"]).value, "Không tìm thấy tài khoản Google")

    def test_update_line_failed_does_not_touch_phase2_columns(self) -> None:
        self._write_workbook(
            ["email", "gift_code", "phase2_status"],
            [["user@example.com", "GRM29G7WLBMZVL3C", "SUCCESS"]],
        )
        store = XlsxStore(self.path)
        store.update_line_failed("user@example.com", "lỗi bất kỳ")

        wb = openpyxl.load_workbook(self.path)
        ws = wb["Mails"]
        headers = {str(c.value): idx + 1 for idx, c in enumerate(ws[1])}
        self.assertEqual(ws.cell(row=2, column=headers["gift_code"]).value, "GRM29G7WLBMZVL3C")
        self.assertEqual(ws.cell(row=2, column=headers["phase2_status"]).value, "SUCCESS")

    def test_update_google_prepare_success_auto_creates_mapping_columns(self) -> None:
        self._write_workbook(["email"], [["user@example.com"]])
        store = XlsxStore(self.path)

        store.update_google_prepare_success("user@example.com", "LineViet-g01")

        wb = openpyxl.load_workbook(self.path)
        ws = wb["Mails"]
        header_row = [cell.value for cell in ws[1]]
        for name in GOOGLE_PREPARE_HEADERS:
            self.assertIn(name, header_row)

        headers = {str(c.value): idx + 1 for idx, c in enumerate(ws[1])}
        self.assertEqual(ws.cell(row=2, column=headers["google_instance_status"]).value, "SUCCESS|LineViet-g01")

    def test_get_pending_accounts_reads_google_prepare_mapping(self) -> None:
        self._write_workbook(
            ["email", "google_instance_status"],
            [["user@example.com", "SUCCESS"]],
        )
        store = XlsxStore(self.path)

        account = store.get_pending_accounts()[0]

        self.assertEqual(account.google_instance_status, "SUCCESS")
        self.assertEqual(account.google_instance, "")
        self.assertEqual(account.google_prepare_status, "SUCCESS")
        self.assertEqual(account.google_prepare_message, "")

    def test_get_pending_accounts_reads_google_prepare_instance_mapping(self) -> None:
        self._write_workbook(
            ["email", "google_instance_status"],
            [["user@example.com", "SUCCESS|LineViet-g01"]],
        )
        store = XlsxStore(self.path)

        account = store.get_pending_accounts()[0]

        self.assertEqual(account.google_prepare_status, "SUCCESS")
        self.assertEqual(account.google_instance, "LineViet-g01")
        self.assertEqual(account.google_prepare_message, "")

    def test_get_pending_accounts_skips_success_rows(self) -> None:
        self._write_workbook(
            ["email", "display_name", "account_password", "line_status"],
            [
                ["done@example.com", "Done User", "pw1", "SUCCESS"],
                ["pending@example.com", "Pending User", "pw2", ""],
                ["failed@example.com", "Failed User", "pw3", "FAILED"],
            ],
        )
        store = XlsxStore(self.path)
        rows = store.get_pending_accounts()
        emails = [row.email for row in rows]
        self.assertEqual(emails, ["pending@example.com", "failed@example.com"])

    def test_get_runnable_accounts_keeps_line_success_when_phase2_failed(self) -> None:
        self._write_workbook(
            ["email", "line_status", "phase2_status", "gift_code", "gift_code_ministop"],
            [
                ["phase2-failed@example.com", "SUCCESS", "FAILED", "", ""],
                ["new@example.com", "", "", "", ""],
            ],
        )
        store = XlsxStore(self.path)

        rows = store.get_runnable_accounts(limit=1)

        self.assertEqual([row.email for row in rows], ["phase2-failed@example.com"])

    def test_get_runnable_accounts_skips_only_when_line_and_phase2_codes_done(self) -> None:
        self._write_workbook(
            ["email", "line_status", "phase2_status", "gift_code", "gift_code_ministop"],
            [
                ["done@example.com", "SUCCESS", "SUCCESS", "AAA111BBB222", "CCC333DDD444"],
                ["next@example.com", "", "", "", ""],
            ],
        )
        store = XlsxStore(self.path)

        rows = store.get_runnable_accounts(limit=1)

        self.assertEqual([row.email for row in rows], ["next@example.com"])

    def test_get_pending_accounts_respects_limit(self) -> None:
        self._write_workbook(
            ["email"],
            [["a@example.com"], ["b@example.com"], ["c@example.com"]],
        )
        store = XlsxStore(self.path)
        rows = store.get_pending_accounts(limit=2)
        self.assertEqual([row.email for row in rows], ["a@example.com", "b@example.com"])

    def test_get_pending_accounts_skips_blank_email_rows(self) -> None:
        self._write_workbook(["email"], [[""], ["user@example.com"]])
        store = XlsxStore(self.path)
        rows = store.get_pending_accounts()
        self.assertEqual([row.email for row in rows], ["user@example.com"])

    def test_get_pending_accounts_never_creates_rows(self) -> None:
        self._write_workbook(["email"], [["user@example.com"]])
        store = XlsxStore(self.path)
        store.get_pending_accounts()
        self.assertEqual(openpyxl.load_workbook(self.path)["Mails"].max_row, 2)

    def test_get_account_returns_none_when_email_missing(self) -> None:
        self._write_workbook(["email"], [["someone@example.com"]])
        store = XlsxStore(self.path)
        self.assertIsNone(store.get_account("nobody@example.com"))

    def test_get_account_returns_row_even_when_already_success(self) -> None:
        #: Khác `get_pending_accounts` (bỏ qua dòng SUCCESS) — `get_account`
        #: phải đọc được để CLI có thể rerun thủ công một email đã xong.
        self._write_workbook(
            ["email", "display_name", "account_password", "line_status"],
            [["user@example.com", "Nguyen Van A", "Secret@123", "SUCCESS"]],
        )
        store = XlsxStore(self.path)
        account = store.get_account("user@example.com")
        self.assertIsNotNone(account)
        self.assertEqual(account.display_name, "Nguyen Van A")
        self.assertEqual(account.account_password, "Secret@123")
        self.assertEqual(account.line_status, "SUCCESS")

    def test_get_account_falls_back_to_name_header(self) -> None:
        #: `_DISPLAY_NAME_HEADERS = ("display_name", "name")` — sheet của
        #: LineViet gốc có thể chỉ có cột `name`.
        self._write_workbook(["email", "name"], [["user@example.com", "Tran Thi B"]])
        store = XlsxStore(self.path)
        account = store.get_account("user@example.com")
        self.assertEqual(account.display_name, "Tran Thi B")

    def test_get_account_matches_pipe_separated_emails_case_insensitively(self) -> None:
        self._write_workbook(["email"], [["First@Example.com|second@example.com"]])
        store = XlsxStore(self.path)
        account = store.get_account("SECOND@EXAMPLE.COM")
        self.assertIsNotNone(account)


if __name__ == "__main__":
    unittest.main()
