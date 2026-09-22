"""Đọc/ghi Excel (sheet "Mails") cho cả Phase 1 (đăng ký LINE) và Phase 2 (gift
code), tương thích cột với LineViet.

Thu hẹp từ `WIN/LineViet/src/connections/xlsx_connection.py::update_email_phase2`
xuống đúng phần Phase 2 cần, và thêm một hành vi KHÔNG có ở bản gốc: khi ghi
`phase2_status=SUCCESS`/`line_status=SUCCESS`, cột `error_details` cũ (nếu
còn) sẽ bị xoá — bản gốc để lại lỗi cũ treo trên dòng dù account đã xong.

Cột đọc/ghi Phase 2 (`PHASE2_HEADERS`) đã xác nhận khớp đúng tên trong file
Excel thật mà `config.json` của project này trỏ tới: `gift_code`,
`gift_code_ministop`, `phase2_status`, `phase2_message`, `phase2_updated_at`.

Cột Phase 1 (`LINE_HEADERS`: `line_status`, `error_details`) tự tạo nếu sheet
chưa có; các cột đầu vào (`email`, `email_password`, `display_name`/`name`,
`account_password`, `status`) PHẢI đã có sẵn trong Excel — không bao giờ tự
tạo dòng email mới.
"""

from __future__ import annotations

import os
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from core.logging_utils import get_logger

_LOG = get_logger("xlsx_store")

PHASE2_HEADERS = ("gift_code", "gift_code_ministop", "phase2_status", "phase2_message", "phase2_updated_at")
LINE_HEADERS = ("line_status", "error_details")
GOOGLE_PREPARE_HEADERS = ("google_instance_status",)
_EMAIL_HEADER = "email"
_EMAIL_PASSWORD_HEADER = "email_password"
_DISPLAY_NAME_HEADERS = ("display_name", "name")
_ACCOUNT_PASSWORD_HEADER = "account_password"
_STATUS_HEADER = "status"
_LINE_STATUS_HEADER = "line_status"
_ERROR_DETAILS_HEADER = "error_details"
_MAX_CELL_LEN = 1500


@dataclass(frozen=True)
class AccountRow:
    """Một dòng account đọc từ Excel — dùng cho `get_pending_accounts`."""

    row: int
    email: str
    email_password: str
    display_name: str
    account_password: str
    status: str
    line_status: str
    error_details: str
    gift_code: str
    gift_code_ministop: str
    phase2_status: str
    google_instance_status: str

    @property
    def google_prepare_status(self) -> str:
        return self.google_instance_status.split("|", 1)[0].strip()

    @property
    def google_instance(self) -> str:
        parts = self.google_instance_status.split("|", 2)
        return parts[1].strip() if len(parts) >= 2 else ""

    @property
    def google_prepare_message(self) -> str:
        parts = self.google_instance_status.split("|", 2)
        return parts[2].strip() if len(parts) >= 3 else ""


class XlsxStoreError(RuntimeError):
    """Không đọc/ghi được file Excel, hoặc không tìm thấy dòng theo email."""


class XlsxStore:
    """Đọc/ghi cột Phase 2 trên sheet "Mails" của file Excel dùng chung với LineViet."""

    def __init__(self, path: str | Path, sheet_name: str = "Mails") -> None:
        self._path = Path(path)
        self._sheet_name = sheet_name
        if not self._path.exists():
            raise XlsxStoreError(f"Không tìm thấy file Excel: {self._path}")

    def _load(self):
        try:
            return openpyxl.load_workbook(self._path)
        except Exception as exc:  # noqa: BLE001 - lỗi đọc file, báo rõ rồi dừng
            raise XlsxStoreError(f"Không đọc được {self._path}: {exc}") from exc

    def _sheet(self, wb) -> Worksheet:
        if self._sheet_name not in wb.sheetnames:
            raise XlsxStoreError(f"Không có sheet `{self._sheet_name}` trong {self._path}.")
        return wb[self._sheet_name]

    @staticmethod
    def _header_map(ws: Worksheet) -> dict[str, int]:
        headers: dict[str, int] = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            name = str(cell.value or "").strip()
            if name:
                headers[name] = col_idx
        return headers

    @staticmethod
    def _ensure_headers(ws: Worksheet, headers: dict[str, int], names: tuple[str, ...]) -> bool:
        changed = False
        next_col = ws.max_column + 1
        for name in names:
            if name not in headers:
                ws.cell(row=1, column=next_col, value=name)
                headers[name] = next_col
                next_col += 1
                changed = True
        return changed

    @staticmethod
    def _find_row(ws: Worksheet, email_col: int, email: str) -> int | None:
        target = email.strip().casefold()
        for row_idx in range(2, ws.max_row + 1):
            raw = ws.cell(row=row_idx, column=email_col).value
            if not raw:
                continue
            candidates = str(raw).split("|")
            if any(candidate.strip().casefold() == target for candidate in candidates):
                return row_idx
        return None

    def find_email(self, email: str) -> bool:
        """True nếu email đã có dòng trong sheet.

        Dùng để preflight TRƯỚC KHI đụng tới thiết bị thật — tránh phí một
        lượt chạy (có thể claim gift code thật) rồi mới phát hiện email sai
        không có trong Excel.
        """
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        return self._find_row(ws, email_col, email) is not None

    def read_existing_codes(self, email: str) -> tuple[str, str]:
        """Đọc `gift_code`/`gift_code_ministop` hiện có của một email (rỗng nếu chưa có)."""
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        row_idx = self._find_row(ws, email_col, email)
        if row_idx is None:
            return "", ""
        gift_col = headers.get("gift_code")
        ministop_col = headers.get("gift_code_ministop")
        gift_code = str(ws.cell(row=row_idx, column=gift_col).value or "") if gift_col else ""
        gift_code_ministop = (
            str(ws.cell(row=row_idx, column=ministop_col).value or "") if ministop_col else ""
        )
        return gift_code.strip(), gift_code_ministop.strip()

    def read_existing_phase2_codes(self, email: str) -> tuple[str, str]:
        """Alias của `read_existing_codes` — tên rõ nghĩa hơn khi gọi từ luồng
        Phase 1 (kiểm tra email đã có gift code từ lần chạy Phase 2 trước hay chưa)."""
        return self.read_existing_codes(email)

    @staticmethod
    def _account_columns(headers: dict[str, int]) -> dict[str, int | None]:
        return {
            "display_name": next((headers[name] for name in _DISPLAY_NAME_HEADERS if name in headers), None),
            "email_password": headers.get(_EMAIL_PASSWORD_HEADER),
            "account_password": headers.get(_ACCOUNT_PASSWORD_HEADER),
            "status": headers.get(_STATUS_HEADER),
            "line_status": headers.get(_LINE_STATUS_HEADER),
            "error_details": headers.get(_ERROR_DETAILS_HEADER),
            "gift_code": headers.get("gift_code"),
            "gift_code_ministop": headers.get("gift_code_ministop"),
            "phase2_status": headers.get("phase2_status"),
            "google_instance_status": headers.get("google_instance_status"),
        }

    @staticmethod
    def _row_to_account(ws: Worksheet, cols: dict[str, int | None], row_idx: int, email: str) -> AccountRow:
        def _cell(col: int | None) -> str:
            if col is None:
                return ""
            return str(ws.cell(row=row_idx, column=col).value or "").strip()

        return AccountRow(
            row=row_idx,
            email=email,
            email_password=_cell(cols["email_password"]),
            display_name=_cell(cols["display_name"]),
            account_password=_cell(cols["account_password"]),
            status=_cell(cols["status"]),
            line_status=_cell(cols["line_status"]),
            error_details=_cell(cols["error_details"]),
            gift_code=_cell(cols["gift_code"]),
            gift_code_ministop=_cell(cols["gift_code_ministop"]),
            phase2_status=_cell(cols["phase2_status"]),
            google_instance_status=_cell(cols["google_instance_status"]),
        )

    def get_pending_accounts(self, limit: int = 0) -> list[AccountRow]:
        """Đọc các dòng CHƯA xong Phase 1 (`line_status` khác `SUCCESS`), theo
        đúng thứ tự trong sheet. `limit<=0` nghĩa là không giới hạn.

        Chỉ ĐỌC — không bao giờ tạo dòng email mới, không ghi gì vào file.
        """
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        cols = self._account_columns(headers)

        rows: list[AccountRow] = []
        for row_idx in range(2, ws.max_row + 1):
            email = str(ws.cell(row=row_idx, column=email_col).value or "").strip()
            if not email:
                continue
            account = self._row_to_account(ws, cols, row_idx, email)
            if account.line_status.upper() == "SUCCESS":
                continue
            rows.append(account)
            if limit > 0 and len(rows) >= limit:
                break
        return rows

    def get_runnable_accounts(self, limit: int = 0) -> list[AccountRow]:
        """Đọc các dòng còn việc phải chạy, theo đúng thứ tự trong sheet.

        Khác `get_pending_accounts` chỉ xét Phase 1: GUI full flow cần pick
        lại cả dòng đã `line_status=SUCCESS` nhưng Phase 2 chưa `SUCCESS` hoặc
        còn thiếu một trong hai mã gift code.
        """
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        cols = self._account_columns(headers)

        rows: list[AccountRow] = []
        for row_idx in range(2, ws.max_row + 1):
            email = str(ws.cell(row=row_idx, column=email_col).value or "").strip()
            if not email:
                continue
            account = self._row_to_account(ws, cols, row_idx, email)
            line_done = account.line_status.strip().upper() == "SUCCESS"
            phase2_done = account.phase2_status.strip().upper() == "SUCCESS"
            has_both_codes = bool(account.gift_code.strip() and account.gift_code_ministop.strip())
            if line_done and phase2_done and has_both_codes:
                continue
            rows.append(account)
            if limit > 0 and len(rows) >= limit:
                break
        return rows

    def get_account(self, email: str) -> AccountRow | None:
        """Đọc MỘT dòng theo email, bất kể `line_status` hiện tại.

        Khác `get_pending_accounts` (bỏ qua dòng đã `SUCCESS`): dùng khi CLI
        `phase1_test.py` cần điền hồ sơ (display_name/account_password) cho
        đúng MỘT email đã biết trước, kể cả khi dòng đó rerun thủ công sau khi
        đã `SUCCESS`/`FAILED`. Trả `None` nếu không có dòng nào khớp email —
        không bao giờ tạo dòng mới.
        """
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        row_idx = self._find_row(ws, email_col, email)
        if row_idx is None:
            return None
        cols = self._account_columns(headers)
        return self._row_to_account(ws, cols, row_idx, email.strip())

    def update_line_success(self, email: str) -> None:
        """Ghi `line_status=SUCCESS` (và `status=SUCCESS` nếu cột đó đã có
        sẵn), xoá `error_details` cũ — không đụng tới cột gift_code/phase2."""
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        self._ensure_headers(ws, headers, LINE_HEADERS)

        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        row_idx = self._find_row(ws, email_col, email)
        if row_idx is None:
            raise XlsxStoreError(f"Không tìm thấy dòng nào khớp email `{email}` trong `{self._path}`.")

        ws.cell(row=row_idx, column=headers[_LINE_STATUS_HEADER], value="SUCCESS")
        status_col = headers.get(_STATUS_HEADER)
        if status_col is not None:
            ws.cell(row=row_idx, column=status_col, value="SUCCESS")
        #: Xem ghi chú ở `update_phase2` — phải gán `.value = None` trực tiếp
        #: mới thực sự xoá được nội dung ô cũ.
        ws.cell(row=row_idx, column=headers[_ERROR_DETAILS_HEADER]).value = None

        self._atomic_save(wb)
        _LOG.info("[%s] Đã ghi Excel: line_status=SUCCESS -> %s", email, self._path)

    def update_line_failed(self, email: str, error: str) -> None:
        """Ghi `line_status=FAILED` + `error_details` — không đụng tới cột
        gift_code/phase2 dù đã có từ lần chạy Phase 2 trước đó."""
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        self._ensure_headers(ws, headers, LINE_HEADERS)

        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        row_idx = self._find_row(ws, email_col, email)
        if row_idx is None:
            raise XlsxStoreError(f"Không tìm thấy dòng nào khớp email `{email}` trong `{self._path}`.")

        ws.cell(row=row_idx, column=headers[_LINE_STATUS_HEADER], value="FAILED")
        status_col = headers.get(_STATUS_HEADER)
        if status_col is not None:
            ws.cell(row=row_idx, column=status_col, value="FAILED")
        ws.cell(row=row_idx, column=headers[_ERROR_DETAILS_HEADER], value=str(error)[:_MAX_CELL_LEN])

        self._atomic_save(wb)
        _LOG.info("[%s] Đã ghi Excel: line_status=FAILED error=%s -> %s", email, error, self._path)

    def update_google_prepare_success(self, email: str, instance_name: str) -> None:
        """Ghi trạng thái Google đã chuẩn bị xong."""
        self._update_google_prepare(
            email,
            instance_name=instance_name,
            status="SUCCESS",
            message="",
        )

    def update_google_prepare_failed(self, email: str, instance_name: str, message: str) -> None:
        """Ghi trạng thái chuẩn bị Google thất bại."""
        self._update_google_prepare(
            email,
            instance_name=instance_name,
            status="FAILED",
            message="",
        )

    def _update_google_prepare(self, email: str, *, instance_name: str, status: str, message: str) -> None:
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        self._ensure_headers(ws, headers, GOOGLE_PREPARE_HEADERS)

        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        row_idx = self._find_row(ws, email_col, email)
        if row_idx is None:
            raise XlsxStoreError(f"Không tìm thấy dòng nào khớp email `{email}` trong `{self._path}`.")

        value = status
        if instance_name:
            value = f"{value}|{instance_name}"
        if message:
            value = f"{value}|{message[:_MAX_CELL_LEN]}"
        ws.cell(row=row_idx, column=headers["google_instance_status"], value=value)

        self._atomic_save(wb)
        _LOG.info(
            "[%s] Đã ghi Excel: google_instance_status=%s -> %s",
            email, status, self._path,
        )

    def update_phase2(
        self,
        email: str,
        *,
        status: str,
        message: str,
        gift_code: str = "",
        gift_code_ministop: str = "",
    ) -> None:
        """Ghi kết quả Phase 2 cho một email, tự tạo cột nếu sheet chưa có.

        Chỉ ghi các giá trị KHÔNG rỗng, để không ghi đè mất giá trị cũ của cột
        không liên quan tới lần chạy này (vd. PARTIAL chỉ có một trong hai mã).
        Riêng khi `status == "SUCCESS"`: xoá `error_details` cũ nếu cột đó có.
        """
        wb = self._load()
        ws = self._sheet(wb)
        headers = self._header_map(ws)
        self._ensure_headers(ws, headers, PHASE2_HEADERS)

        email_col = headers.get(_EMAIL_HEADER)
        if email_col is None:
            raise XlsxStoreError(f"Sheet `{self._sheet_name}` thiếu cột `{_EMAIL_HEADER}`.")
        row_idx = self._find_row(ws, email_col, email)
        if row_idx is None:
            raise XlsxStoreError(f"Không tìm thấy dòng nào khớp email `{email}` trong `{self._path}`.")

        values = {
            "gift_code": gift_code,
            "gift_code_ministop": gift_code_ministop,
            "phase2_status": status,
            "phase2_message": message,
            "phase2_updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        for name, value in values.items():
            if value in (None, ""):
                continue
            col = headers[name]
            ws.cell(row=row_idx, column=col, value=str(value)[:_MAX_CELL_LEN])

        if status == "SUCCESS":
            error_col = headers.get(_ERROR_DETAILS_HEADER)
            if error_col is not None:
                #: `ws.cell(..., value=None)` KHÔNG xoá được ô — openpyxl coi
                #: `value=None` là "không truyền gì" (tham số mặc định) nên bỏ
                #: qua, không gán. Phải lấy Cell rồi gán `.value = None` trực
                #: tiếp mới thực sự xoá nội dung cũ.
                ws.cell(row=row_idx, column=error_col).value = None

        self._atomic_save(wb)
        _LOG.info(
            "[%s] Đã ghi Excel: status=%s gift_code=%s gift_code_ministop=%s -> %s",
            email, status, gift_code or "-", gift_code_ministop or "-", self._path,
        )

    def _atomic_save(self, wb) -> None:
        tmp_path = self._path.with_name(f"{self._path.stem}.tmp_{int(time.time() * 1000)}{self._path.suffix}")
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, self._path)
        except OSError as exc:
            tmp_path.unlink(missing_ok=True)
            raise XlsxStoreError(
                f"Không ghi được {self._path} (file có thể đang mở trong Excel hoặc bị khoá): {exc}"
            ) from exc


__all__ = [
    "AccountRow",
    "GOOGLE_PREPARE_HEADERS",
    "LINE_HEADERS",
    "PHASE2_HEADERS",
    "XlsxStore",
    "XlsxStoreError",
]
